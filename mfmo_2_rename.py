# Fotos_Videos_Sortieren_Rename
# Erweitert um Time_Shift-Logik und Ziel_Name_Time_Shift

import os
import sys
import logging
import configparser
import pandas as pd
import string
from datetime import datetime, timedelta
from tkinter import Tk, filedialog
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

def set_arbeitsverzeichnis():
    skriptpfad = os.path.abspath(sys.argv[0])
    os.chdir(os.path.dirname(skriptpfad))

def screen_and_log(message):
    print(message)
    logging.info(message)

def waehle_excel_datei(startpfad):
    root = Tk()
    root.withdraw()
    datei = filedialog.askopenfilename(
        initialdir=startpfad,
        title="Bitte Excel-Datei zur Analyse auswählen",
        filetypes=[("Excel-Dateien", "*.xlsx")]
    )
    root.destroy()
    return os.path.normpath(datei) if datei else None

def lade_konfiguration(pfad='mfmo_2_rename.ini'):
    if not os.path.isfile(pfad):
        screen_and_log(f"INI-Datei nicht gefunden: {pfad}")
        raise SystemExit(1)

    config = configparser.ConfigParser()
    config.read(pfad, encoding='utf-8')

    try:
        quelle_pfad = os.path.normpath(config.get('ALLGEMEIN', 'quelle_pfad'))
        log_datei = config.get('ALLGEMEIN', 'log_datei')
        exiftool_pfad = os.path.normpath(config.get('ALLGEMEIN', 'exiftool_pfad'))
        import_excel_pfad = config.get('ALLGEMEIN', 'import_excel_pfad')
        return quelle_pfad, log_datei, exiftool_pfad, import_excel_pfad
    except Exception as e:
        screen_and_log(f"Fehler beim Einlesen der INI-Datei: {e}")
        raise SystemExit(1)

def initialisiere_logging(log_datei):
    logging.basicConfig(filename=log_datei, level=logging.INFO,
                        format='%(asctime)s - %(levelname)s - %(message)s',
                        encoding='utf-8')

def berechne_zielname_shift(df):
    df['Time_Shift'] = pd.to_numeric(df.get('Time_Shift', 0), errors='coerce').fillna(0).astype(int)
    ziel_shift_namen = []

    for i, row in df.iterrows():
        ziel = row.get("Ziel_Name")
        if pd.isna(ziel):
            ziel_shift_namen.append(None)
            continue
        try:
            basename, ext = os.path.splitext(ziel)
            if '_' not in basename:
                ziel_shift_namen.append(None)
                continue
            datum, uhrzeit = basename.split('_')
            dt = datetime.strptime(datum + uhrzeit, "%Y%m%d%H%M%S")
            dt_shift = dt + timedelta(hours=row['Time_Shift'])
            ziel_shift_namen.append(dt_shift.strftime("%Y%m%d_%H%M%S") + ext)
        except Exception as e:
            ziel_shift_namen.append(None)

    df['Ziel_Name_Time_Shift'] = ziel_shift_namen
    return df

def bereinige_dubletten(df):
    counter = {}
    neue_namen = []

    for zielname in df['Ziel_Name_Time_Shift']:
        
        if pd.isna(zielname):
            neue_namen.append(None)
            continue

        basisname, ext = os.path.splitext(zielname)
        key = zielname.lower()
        counter[key] = counter.get(key, 0) + 1
        buchstabe = string.ascii_lowercase[counter[key] - 1]
        neuer_name = f"{basisname}{buchstabe}{ext}"
        neue_namen.append(neuer_name)

    df['Ziel_Name_bereinigt'] = neue_namen
    return df

def format_excel_tabelle(dateipfad):
    try:
        wb = load_workbook(dateipfad)
        ws = wb.active
        ws.freeze_panes = "A2"
        max_row = ws.max_row
        max_col = ws.max_column
        tabellenbereich = f"A1:{chr(64 + max_col)}{max_row}"
        tabelle = Table(displayName="BereinigtTabelle", ref=tabellenbereich)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tabelle.tableStyleInfo = style
        ws.add_table(tabelle)
        for column_cells in ws.columns:
            max_length = 0
            col_letter = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = max_length + 2
        wb.save(dateipfad)
        screen_and_log("Excel-Tabelle formatiert, Kopfzeile fixiert und Spaltenbreiten angepasst.")
    except Exception as e:
        screen_and_log(f"Fehler bei der Formatierung der Excel-Tabelle: {e}")

def main():
    set_arbeitsverzeichnis()
    quelle_pfad, log_datei, exiftool_pfad, import_excel_pfad_vorauswahl = lade_konfiguration()
    initialisiere_logging(log_datei)
    screen_and_log("Toolstart: Ziel_Name_Time_Shift + Dublettenprüfung")

    ausgewaehlte_excel = waehle_excel_datei(os.path.dirname(import_excel_pfad_vorauswahl))
    if not ausgewaehlte_excel:
        screen_and_log("Keine Datei ausgewählt. Vorgang abgebrochen.")
        return

    ziel_excel = ausgewaehlte_excel.replace("_analyse.xlsx", "_rename.xlsx")

    try:
        df = pd.read_excel(ausgewaehlte_excel)
        if 'Ziel_Name' not in df.columns:
            screen_and_log("Spalte 'Ziel_Name' nicht gefunden.")
            return

        df = berechne_zielname_shift(df)
        df = bereinige_dubletten(df)
        df.to_excel(ziel_excel, index=False)
        format_excel_tabelle(ziel_excel)
        screen_and_log(f"Datei gespeichert unter: {ziel_excel}")

    except Exception as e:
        screen_and_log(f"Fehler bei der Verarbeitung: {e}")

    screen_and_log("Toolende")

if __name__ == "__main__":
    main()
