# Media From Mobile Organize - Combined Workflow
# Combines analyze, rename, and execute steps with pause points for user editing

import os
import sys
import logging
import configparser
import pandas as pd
import subprocess
import json
import re
import string
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from tqdm import tqdm
from tkinter import Tk, filedialog

# ============================================================================
# UTILITY FUNCTIONS
# ============================================================================

def set_arbeitsverzeichnis():
    """Set working directory to script location"""
    skriptpfad = os.path.abspath(sys.argv[0])
    os.chdir(os.path.dirname(skriptpfad))

def screen_and_log(message):
    """Print to screen and write to log"""
    print(message)
    logging.info(message)

def initialisiere_logging(log_datei):
    """Initialize logging"""
    logging.basicConfig(
        filename=log_datei,
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        encoding='utf-8'
    )

def waehle_verzeichnis(startpfad):
    """Show dialog to select a folder"""
    root = Tk()
    root.withdraw()
    ausgewaehlt = filedialog.askdirectory(
        initialdir=startpfad,
        title="Bitte Quellordner wählen"
    )
    root.destroy()
    return os.path.normpath(ausgewaehlt) if ausgewaehlt else startpfad

def ist_datei_geschlossen(dateiname):
    """Check if a file is closed by trying to open it in write mode"""
    try:
        # Try to open the file in append mode
        with open(dateiname, 'a'):
            pass
        return True
    except IOError:
        return False

def pause_fuer_bearbeitung(dateiname):
    """Pause execution and wait for user to edit the file"""
    screen_and_log("\n" + "="*60)
    screen_and_log(f"Excel-Datei erstellt: {dateiname}")
    screen_and_log("Bitte bearbeiten Sie jetzt die Excel-Datei.")
    screen_and_log("Check for timeshift per user, and time difference for iPhone video vs picture (1h in winter, 2h in summer).")
    screen_and_log("Add names for files without identified metadata.")
    screen_and_log("="*60)

    while True:
        input("\nDrücken Sie ENTER, wenn Sie fertig sind und fortfahren möchten...")

        # Check if file is closed
        if ist_datei_geschlossen(dateiname):
            screen_and_log("Excel-Datei ist geschlossen. Fortfahren...")
            break
        else:
            screen_and_log("\n" + "!"*60)
            screen_and_log("WARNUNG: Die Excel-Datei ist noch geöffnet!")
            screen_and_log("Bitte schließen Sie die Datei in Excel und speichern Sie Ihre Änderungen.")
            screen_and_log("!"*60)

def lade_konfiguration(pfad='mfmo.ini'):
    """Load configuration from INI file"""
    if not os.path.isfile(pfad):
        screen_and_log(f"INI-Datei nicht gefunden: {pfad}")
        raise SystemExit(1)

    config = configparser.ConfigParser()
    config.read(pfad, encoding='utf-8')

    try:
        quelle_pfad = os.path.normpath(config.get('ALLGEMEIN', 'quelle_pfad'))
        log_datei = config.get('ALLGEMEIN', 'log_datei')
        exiftool_pfad = os.path.normpath(config.get('ALLGEMEIN', 'exiftool_pfad'))
        artist_map = {
            k.strip().lower(): v.strip()
            for k, v in config.items('Artist')
        } if config.has_section('Artist') else {}
        return quelle_pfad, log_datei, exiftool_pfad, artist_map
    except Exception as e:
        screen_and_log(f"Fehler beim Einlesen der INI-Datei: {e}")
        raise SystemExit(1)

# ============================================================================
# STEP 1: ANALYZE - Extract metadata and create initial Excel
# ============================================================================

def exiftool_auslesen(dateipfad, exiftool_pfad, artist_map=None):
    """Read EXIF data using ExifTool"""
    try:
        result = subprocess.run(
            [exiftool_pfad, '-j', dateipfad],
            capture_output=True, text=True, check=True
        )
        daten = json.loads(result.stdout)[0]
        modell = daten.get("Model", "").strip().lower()
        kuerzel = artist_map.get(modell, "") if artist_map else ""

        ext = os.path.splitext(dateipfad)[1].lower()
        dateiname = os.path.basename(dateipfad)
        aufnahmedatum = daten.get("DateTimeOriginal") or daten.get("CreateDate")
        ziel_name = None

        # Try to extract date from filename
        match = re.search(r'(\d{8})_(\d{6})', dateiname)
        if match:
            ziel_name = f"{match.group(1)}_{match.group(2)}{ext}"
        elif aufnahmedatum:
            try:
                dt = aufnahmedatum.split()[0].replace(":", "") + "_" + aufnahmedatum.split()[1].replace(":", "")
                ziel_name = dt + ext
            except Exception as e:
                logging.warning(f"Fehler beim Erzeugen von Ziel_Name für {dateipfad}: {e}")

        if not ziel_name:
            logging.warning(f"Kein Ziel_Name erzeugt für Datei: {dateiname}")

        return {
            "Dateiname": dateiname,
            "Endung": ext,
            "EXIF-Aufnahmedatum": aufnahmedatum,
            "Kamera-Fabrikat": daten.get("Make"),
            "Kamera-Modell": modell,
            "GPS-Breite": daten.get("GPSLatitude"),
            "GPS-Länge": daten.get("GPSLongitude"),
            "Artist": kuerzel,
            "Ziel_Name": ziel_name
        }
    except Exception as e:
        logging.warning(f"ExifTool-Fehler bei {dateipfad}: {e}")
        return {
            "Dateiname": os.path.basename(dateipfad),
            "Endung": os.path.splitext(dateipfad)[1].lower(),
            "EXIF-Aufnahmedatum": None,
            "Kamera-Fabrikat": None,
            "Kamera-Modell": None,
            "GPS-Breite": None,
            "GPS-Länge": None,
            "Artist": None,
            "Ziel_Name": None
        }

def verzeichnis_durchsuchen(quelle_pfad, exiftool_pfad, artist_map):
    """Search directory and extract metadata with progress bar"""
    daten = []
    dateien = [f for f in os.listdir(quelle_pfad) if os.path.isfile(os.path.join(quelle_pfad, f))]

    screen_and_log(f"{len(dateien)} Dateien gefunden. Verarbeitung startet...")

    for eintrag in tqdm(dateien, desc="Verarbeite", unit="Datei"):
        dateipfad = os.path.join(quelle_pfad, eintrag)
        info = exiftool_auslesen(dateipfad, exiftool_pfad, artist_map)
        daten.append(info)

    return daten

def format_excel_tabelle(dateiname, table_name="DateiTabelle"):
    """Format Excel file with table style and frozen header"""
    wb = load_workbook(dateiname)
    ws = wb.active

    # Freeze top row
    ws.freeze_panes = "A2"

    # Create table
    max_row = ws.max_row
    max_col = ws.max_column
    tabellenbereich = f"A1:{chr(64 + max_col)}{max_row}"

    tabelle = Table(displayName=table_name, ref=tabellenbereich)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    tabelle.tableStyleInfo = style
    ws.add_table(tabelle)

    # Auto-adjust column widths
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[column].width = adjusted_width

    wb.save(dateiname)

def step_1_analyze(quelle_pfad, exiftool_pfad, artist_map):
    """Step 1: Analyze media files and create Excel"""
    screen_and_log("\n" + "="*60)
    screen_and_log("SCHRITT 1: ANALYSE DER MEDIENDATEIEN")
    screen_and_log("="*60)

    ordnername = os.path.basename(os.path.normpath(quelle_pfad))
    export_excel_pfad = os.path.join(quelle_pfad, f"{ordnername}_analyse.xlsx")

    daten = verzeichnis_durchsuchen(quelle_pfad, exiftool_pfad, artist_map)

    # Create DataFrame and add Time_Shift column
    df = pd.DataFrame(daten)
    if "Ziel_Name" in df.columns:
        ziel_index = df.columns.get_loc("Ziel_Name")
        df.insert(ziel_index, "Time_Shift", "")
    else:
        df["Time_Shift"] = ""

    df.to_excel(export_excel_pfad, index=False)
    format_excel_tabelle(export_excel_pfad, "DateiTabelle")

    screen_and_log(f"Analyse abgeschlossen. Excel-Datei erstellt: {export_excel_pfad}")

    return export_excel_pfad

# ============================================================================
# STEP 2: RENAME - Process time shifts and handle duplicates
# ============================================================================

def berechne_zielname_shift(df):
    """Calculate target name with time shift applied"""
    df['Time_Shift'] = pd.to_numeric(df.get('Time_Shift', 0), errors='coerce').fillna(0).astype(int)
    ziel_shift_namen = []

    for i, row in df.iterrows():
        ziel = row.get("Ziel_Name")
        if pd.isna(ziel):
            ziel_shift_namen.append(None)
            continue
        try:
            basename, ext = os.path.splitext(ziel)
            if '_' not in basename:
                ziel_shift_namen.append(None)
                continue
            datum, uhrzeit = basename.split('_')
            dt = datetime.strptime(datum + uhrzeit, "%Y%m%d%H%M%S")
            dt_shift = dt + timedelta(hours=row['Time_Shift'])
            ziel_shift_namen.append(dt_shift.strftime("%Y%m%d_%H%M%S") + ext)
        except Exception as e:
            ziel_shift_namen.append(None)

    df['Ziel_Name_Time_Shift'] = ziel_shift_namen
    return df

def bereinige_dubletten(df):
    """Clean duplicates by adding letters (a, b, c, ...)"""
    counter = {}
    neue_namen = []

    for zielname in df['Ziel_Name_Time_Shift']:
        if pd.isna(zielname):
            neue_namen.append(None)
            continue

        basisname, ext = os.path.splitext(zielname)
        key = zielname.lower()
        counter[key] = counter.get(key, 0) + 1
        buchstabe = string.ascii_lowercase[counter[key] - 1]
        neuer_name = f"{basisname}{buchstabe}{ext}"
        neue_namen.append(neuer_name)

    df['Ziel_Name_bereinigt'] = neue_namen
    return df

def step_2_rename(analyse_excel_pfad):
    """Step 2: Process time shifts and clean duplicates"""
    screen_and_log("\n" + "="*60)
    screen_and_log("SCHRITT 2: TIMESHIFT UND DUBLETTENPRÜFUNG")
    screen_and_log("="*60)

    ziel_excel = analyse_excel_pfad.replace("_analyse.xlsx", "_rename.xlsx")

    try:
        df = pd.read_excel(analyse_excel_pfad)
        if 'Ziel_Name' not in df.columns:
            screen_and_log("FEHLER: Spalte 'Ziel_Name' nicht gefunden.")
            return None

        df = berechne_zielname_shift(df)
        df = bereinige_dubletten(df)
        df.to_excel(ziel_excel, index=False)
        format_excel_tabelle(ziel_excel, "BereinigtTabelle")

        screen_and_log(f"Verarbeitung abgeschlossen. Excel-Datei erstellt: {ziel_excel}")
        return ziel_excel

    except Exception as e:
        screen_and_log(f"FEHLER bei der Verarbeitung: {e}")
        return None

# ============================================================================
# STEP 3: EXECUTE - Perform actual file renaming
# ============================================================================

def fuehre_umbenennung_aus(excel_pfad, ordnerpfad):
    """Execute file renaming based on Excel data"""
    try:
        df = pd.read_excel(excel_pfad)
    except Exception as e:
        screen_and_log(f"Fehler beim Einlesen der Excel-Datei: {e}")
        return

    if 'Dateiname' not in df.columns or 'Ziel_Name_bereinigt' not in df.columns:
        screen_and_log("FEHLER: Spalten 'Dateiname' oder 'Ziel_Name_bereinigt' fehlen.")
        return

    umbenannt = 0
    fehlerhaft = 0

    for _, row in df.iterrows():
        alt = row['Dateiname']
        neu = row['Ziel_Name_bereinigt']

        if pd.isna(alt) or pd.isna(neu):
            continue

        quellpfad = os.path.join(ordnerpfad, alt)
        zielpfad = os.path.join(ordnerpfad, neu)

        if not os.path.exists(quellpfad):
            screen_and_log(f"FEHLT: Quelldatei nicht gefunden: {quellpfad}")
            fehlerhaft += 1
            continue

        if os.path.exists(zielpfad):
            screen_and_log(f"WARNUNG: Ziel existiert bereits: {zielpfad} - übersprungen")
            fehlerhaft += 1
            continue

        try:
            os.rename(quellpfad, zielpfad)
            screen_and_log(f"UMBENANNT: {alt} -> {neu}")
            umbenannt += 1
        except Exception as e:
            screen_and_log(f"FEHLER bei {alt}: {e}")
            fehlerhaft += 1

    screen_and_log(f"\nZusammenfassung: {umbenannt} Dateien umbenannt, {fehlerhaft} Fehler.")

def step_3_execute(rename_excel_pfad):
    """Step 3: Execute the file renaming"""
    screen_and_log("\n" + "="*60)
    screen_and_log("SCHRITT 3: DATEIEN UMBENENNEN")
    screen_and_log("="*60)

    ordnerpfad = os.path.dirname(rename_excel_pfad)
    fuehre_umbenennung_aus(rename_excel_pfad, ordnerpfad)

# ============================================================================
# MAIN WORKFLOW
# ============================================================================

def main():
    try:
        set_arbeitsverzeichnis()

        # Load configuration
        screen_and_log("Lade Konfiguration...")
        quelle_pfad_ini, log_datei, exiftool_pfad, artist_map = lade_konfiguration()
        initialisiere_logging(log_datei)

        screen_and_log("\n" + "="*60)
        screen_and_log("MEDIA FROM MOBILE ORGANIZE - WORKFLOW START")
        screen_and_log("="*60)

        # Select folder
        quelle_pfad = waehle_verzeichnis(quelle_pfad_ini)
        screen_and_log(f"Ausgewähltes Verzeichnis: {quelle_pfad}")

        # STEP 1: Analyze
        analyse_excel = step_1_analyze(quelle_pfad, exiftool_pfad, artist_map)
        pause_fuer_bearbeitung(analyse_excel)

        # STEP 2: Rename
        rename_excel = step_2_rename(analyse_excel)
        if rename_excel is None:
            screen_and_log("Workflow abgebrochen aufgrund von Fehlern in Schritt 2.")
            return
        pause_fuer_bearbeitung(rename_excel)

        # STEP 3: Execute
        step_3_execute(rename_excel)

        screen_and_log("\n" + "="*60)
        screen_and_log("WORKFLOW ERFOLGREICH ABGESCHLOSSEN")
        screen_and_log("="*60)

    except Exception as e:
        screen_and_log(f"FEHLER im Hauptprogramm: {e}")

if __name__ == "__main__":
    main()
