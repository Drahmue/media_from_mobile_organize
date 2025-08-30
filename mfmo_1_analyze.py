# Fotos_Videos_Sortieren
# Liest Mediendateien aus einem Verzeichnis, extrahiert Metadaten via ExifTool und erzeugt eine Excel-Tabelle

import os
import sys
import logging
import configparser
import pandas as pd
import subprocess
import json
import re
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from tqdm import tqdm
from tkinter import Tk, filedialog


def set_arbeitsverzeichnis():
    skriptpfad = os.path.abspath(sys.argv[0])
    os.chdir(os.path.dirname(skriptpfad))

def screen_and_log(message):
    print(message)
    logging.info(message)

def waehle_verzeichnis(startpfad):
    """Zeigt einen Dialog zur Auswahl eines Ordners. Gibt den gewählten Pfad zurück oder den Startpfad bei Abbruch."""
    root = Tk()
    root.withdraw()
    ausgewaehlt = filedialog.askdirectory(initialdir=startpfad, title="Bitte Quellordner wählen")
    root.destroy()
    return os.path.normpath(ausgewaehlt) if ausgewaehlt else startpfad

def lade_konfiguration(pfad='mfmo_1_analyze.ini'):
    if not os.path.isfile(pfad):
        screen_and_log(f"INI-Datei nicht gefunden: {pfad}")
        raise SystemExit(1)

    config = configparser.ConfigParser()
    config.read(pfad, encoding='utf-8')

    try:
        quelle_pfad = os.path.normpath(config.get('ALLGEMEIN', 'quelle_pfad'))
        log_datei = config.get('ALLGEMEIN', 'log_datei')
        exiftool_pfad = os.path.normpath(config.get('ALLGEMEIN', 'exiftool_pfad'))
        export_excel_pfad = config.get('ALLGEMEIN', 'export_excel_pfad')
        artist_map = {
            k.strip().lower(): v.strip()
            for k, v in config.items('Artist')
        } if config.has_section('Artist') else {}
        return quelle_pfad, log_datei, exiftool_pfad, artist_map, export_excel_pfad
    except Exception as e:
        screen_and_log(f"Fehler beim Einlesen der INI-Datei: {e}")
        raise SystemExit(1)

def initialisiere_logging(log_datei):
    logging.basicConfig(filename=log_datei, level=logging.INFO,
                        format='%(asctime)s - %(levelname)s - %(message)s',
                        encoding='utf-8')


def exiftool_auslesen(dateipfad, exiftool_pfad, artist_map=None):
    try:
        result = subprocess.run(
            [exiftool_pfad, '-j', dateipfad],
            capture_output=True, text=True, check=True
        )
        daten = json.loads(result.stdout)[0]
        modell = daten.get("Model", "").strip().lower()
        kuerzel = artist_map.get(modell, "") if artist_map else ""

        ext = os.path.splitext(dateipfad)[1].lower()
        dateiname = os.path.basename(dateipfad)
        aufnahmedatum = daten.get("DateTimeOriginal") or daten.get("CreateDate")
        ziel_name = None

        match = re.search(r'(\d{8})_(\d{6})', dateiname)
        if match:
            ziel_name = f"{match.group(1)}_{match.group(2)}{ext}"
        elif aufnahmedatum:
            try:
                dt = aufnahmedatum.split()[0].replace(":", "") + "_" + aufnahmedatum.split()[1].replace(":", "")
                ziel_name = dt + ext
            except Exception as e:
                logging.warning(f"Fehler beim Erzeugen von Ziel_Name für {dateipfad}: {e}")

        if not ziel_name:
            logging.warning(f"Kein Ziel_Name erzeugt für Datei: {dateiname}")

        return {
            "Dateiname": dateiname,
            "Endung": ext,
            "EXIF-Aufnahmedatum": aufnahmedatum,
            "Kamera-Fabrikat": daten.get("Make"),
            "Kamera-Modell": modell,
            "GPS-Breite": daten.get("GPSLatitude"),
            "GPS-Länge": daten.get("GPSLongitude"),
            "Artist": kuerzel,
            "Ziel_Name": ziel_name
        }
    except Exception as e:
        logging.warning(f"ExifTool-Fehler bei {dateipfad}: {e}")
        return {
            "Dateiname": os.path.basename(dateipfad),
            "Endung": os.path.splitext(dateipfad)[1].lower(),
            "EXIF-Aufnahmedatum": None,
            "Kamera-Fabrikat": None,
            "Kamera-Modell": None,
            "GPS-Breite": None,
            "GPS-Länge": None,
            "Artist": None,
            "Ziel_Name": None
        }

def verzeichnis_durchsuchen(quelle_pfad, exiftool_pfad, artist_map):
    """Durchsucht ein Verzeichnis und ruft ExifTool auf – mit Fortschrittsbalken"""
    daten = []
    dateien = [f for f in os.listdir(quelle_pfad) if os.path.isfile(os.path.join(quelle_pfad, f))]

    screen_and_log(f"{len(dateien)} Dateien gefunden. Verarbeitung startet...")

    for eintrag in tqdm(dateien, desc="Verarbeite", unit="Datei"):
        dateipfad = os.path.join(quelle_pfad, eintrag)
        info = exiftool_auslesen(dateipfad, exiftool_pfad, artist_map)
        daten.append(info)

    return daten

def speichere_excel(daten, dateiname):
    """Speichert die Dateiinformationen in einer Excel-Tabelle mit Formatierung und fixierter Kopfzeile"""
    df = pd.DataFrame(daten)

    # ➕ Leere "Time_Shift"-Spalte vor "Ziel_Name" einfügen
    if "Ziel_Name" in df.columns:
        ziel_index = df.columns.get_loc("Ziel_Name")
        df.insert(ziel_index, "Time_Shift", "")
    else:
        # Falls "Ziel_Name" fehlt, einfach am Ende anhängen
        df["Time_Shift"] = ""

    df.to_excel(dateiname, index=False)

    wb = load_workbook(dateiname)
    ws = wb.active

    # 🔒 Fixiere die oberste Zeile
    ws.freeze_panes = "A2"

    # Tabellenbereich bestimmen
    max_row = ws.max_row
    max_col = ws.max_column
    tabellenbereich = f"A1:{chr(64 + max_col)}{max_row}"  # z. B. A1:G42

    # Tabelle erstellen
    tabelle = Table(displayName="DateiTabelle", ref=tabellenbereich)
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tabelle.tableStyleInfo = style
    ws.add_table(tabelle)

    # 📏 Spaltenbreite automatisch anpassen
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[column].width = adjusted_width

    wb.save(dateiname)
    screen_and_log(f"Excel-Datei als formatierte Tabelle mit fixierter Kopfzeile gespeichert: {dateiname}")



def main():
    try:
        set_arbeitsverzeichnis()
        screen_and_log("Lade Konfiguration")
        quelle_pfad_ini, log_datei, exiftool_pfad, artist_map, _ = lade_konfiguration()
        initialisiere_logging(log_datei)
        screen_and_log("Programmstart")

        # 📁 Pfadauswahl mit Vorauswahl aus INI
        quelle_pfad = waehle_verzeichnis(quelle_pfad_ini)
        screen_and_log(f"Verzeichnis wird eingelesen: {quelle_pfad}")

        # 📝 Excel-Dateiname automatisch aus Ordnername erzeugen
        ordnername = os.path.basename(os.path.normpath(quelle_pfad))
        export_excel_pfad = os.path.join(quelle_pfad, f"{ordnername}_analyse.xlsx")

        daten = verzeichnis_durchsuchen(quelle_pfad, exiftool_pfad, artist_map)
        speichere_excel(daten, export_excel_pfad)
        screen_and_log("Programm erfolgreich beendet.")
    except Exception as e:
        screen_and_log(f"Fehler im Hauptprogramm: {e}")

if __name__ == "__main__":
    main()
